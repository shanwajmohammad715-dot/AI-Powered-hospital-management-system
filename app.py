import os
import sqlite3
from datetime import datetime
from functools import wraps
from io import BytesIO

from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, session, url_for, send_file
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook

load_dotenv()
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-only-change-this-secret-key")
DB_PATH = os.path.join(os.path.dirname(__file__), "hospital.db")

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def has_column(con, table, column):
    return column in [r["name"] for r in con.execute(f"PRAGMA table_info({table})").fetchall()]

def init_db():
    with db() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS patients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, age INTEGER NOT NULL, gender TEXT NOT NULL,
            phone TEXT NOT NULL, email TEXT DEFAULT '', address TEXT DEFAULT '',
            created_at TEXT NOT NULL, user_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS doctors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, specialty TEXT NOT NULL, fee REAL NOT NULL
        );
        CREATE TABLE IF NOT EXISTS appointments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_id INTEGER NOT NULL REFERENCES patients(id),
            doctor_id INTEGER NOT NULL REFERENCES doctors(id),
            appointment_date TEXT NOT NULL, reason TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Booked', created_at TEXT NOT NULL,
            user_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS bills (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_id INTEGER NOT NULL REFERENCES patients(id),
            description TEXT NOT NULL, amount REAL NOT NULL,
            payment_status TEXT NOT NULL DEFAULT 'Pending',
            created_at TEXT NOT NULL, user_id INTEGER
        );
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE COLLATE NOCASE,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('admin','doctor')),
            doctor_id INTEGER UNIQUE REFERENCES doctors(id),
            created_at TEXT NOT NULL
        );
        """)
        # Safe migration for an existing copy of the original project.
        for table, column in (("patients", "user_id"), ("appointments", "user_id"), ("bills", "user_id")):
            if not has_column(con, table, column):
                con.execute(f"ALTER TABLE {table} ADD COLUMN {column} INTEGER")

        if con.execute("SELECT COUNT(*) FROM doctors").fetchone()[0] == 0:
            con.executemany(
                "INSERT INTO doctors(name,specialty,fee) VALUES(?,?,?)",
                [
                    ("Dr. Ayesha Khan", "General Physician", 400),
                    ("Dr. Rahul Verma", "Cardiology", 800),
                    ("Dr. Neha Singh", "Dermatology", 600),
                    ("Dr. Imran Ali", "Pediatrics", 500),
                ],
            )

        # Seed the admin account once. Change these values in .env before use.
        admin_username = os.getenv("ADMIN_USERNAME", "admin").strip()
        admin_password = os.getenv("ADMIN_PASSWORD", "ChangeAdmin123!")
        existing_admin = con.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()
        if not existing_admin:
            con.execute(
                "INSERT INTO users(username,password_hash,full_name,role,created_at) VALUES(?,?,?,?,?)",
                (admin_username, generate_password_hash(admin_password), "Hospital Administrator", "admin",
                 datetime.now().isoformat(timespec="minutes"))
            )

def current_user():
    return session.get("user_id")

def doctor_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not current_user():
            flash("Please sign in to access protected hospital records.", "warning")
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not current_user():
            flash("Please sign in first.", "warning")
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("Only the administrator can manage doctor accounts.", "danger")
            return redirect(url_for("index"))
        return fn(*args, **kwargs)
    return wrapped

def owned_patient_filter(alias="p"):
    if session.get("role") == "admin":
        return "", []
    return f" AND {alias}.user_id = ?", [current_user()]

@app.route("/")
def index():
    if session.get("role") == "admin":
        with db() as con:
            stats = {
                "patients": con.execute("SELECT COUNT(*) FROM patients").fetchone()[0],
                "appointments": con.execute("SELECT COUNT(*) FROM appointments").fetchone()[0],
                "pending_bills": con.execute("SELECT COUNT(*) FROM bills WHERE payment_status='Pending'").fetchone()[0],
                "doctors": con.execute("SELECT COUNT(*) FROM doctors").fetchone()[0],
            }
            appointments = con.execute("""SELECT a.id,p.name patient_name,d.name doctor_name,d.specialty,
                a.appointment_date,a.status FROM appointments a JOIN patients p ON p.id=a.patient_id
                JOIN doctors d ON d.id=a.doctor_id ORDER BY a.appointment_date DESC LIMIT 6""").fetchall()
    elif current_user():
        uid, did = current_user(), session.get("doctor_id")
        with db() as con:
            stats = {
                "patients": con.execute("SELECT COUNT(*) FROM patients WHERE user_id=?", (uid,)).fetchone()[0],
                "appointments": con.execute("SELECT COUNT(*) FROM appointments WHERE user_id=?", (uid,)).fetchone()[0],
                "pending_bills": con.execute("SELECT COUNT(*) FROM bills WHERE user_id=? AND payment_status='Pending'", (uid,)).fetchone()[0],
                "doctors": con.execute("SELECT COUNT(*) FROM doctors").fetchone()[0],
            }
            appointments = con.execute("""SELECT a.id,p.name patient_name,d.name doctor_name,d.specialty,
                a.appointment_date,a.status FROM appointments a JOIN patients p ON p.id=a.patient_id
                JOIN doctors d ON d.id=a.doctor_id WHERE a.user_id=? ORDER BY a.appointment_date DESC LIMIT 6""",
                (uid,)).fetchall()
    else:
        with db() as con:
            stats = {"patients": 0, "appointments": 0, "pending_bills": 0,
                     "doctors": con.execute("SELECT COUNT(*) FROM doctors").fetchone()[0]}
        appointments = []
    return render_template("index.html", stats=stats, appointments=appointments)

@app.route("/doctors")
def doctors():
    with db() as con:
        rows = con.execute("SELECT * FROM doctors ORDER BY specialty, name").fetchall()
    return render_template("doctors.html", doctors=rows)

@app.route("/admin/doctors", methods=["GET", "POST"])
@admin_required
def manage_doctors():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        try:
            doctor_id = int(request.form.get("doctor_id", ""))
        except ValueError:
            doctor_id = -1
        if len(username) < 3 or not full_name or len(password) < 10 or doctor_id < 1:
            flash("Enter a username (at least 3 characters), full name, existing doctor, and password of at least 10 characters.", "danger")
        else:
            with db() as con:
                if con.execute("SELECT 1 FROM users WHERE username=? COLLATE NOCASE", (username,)).fetchone():
                    flash("That username is already taken. Choose another.", "danger")
                elif con.execute("SELECT 1 FROM users WHERE doctor_id=?", (doctor_id,)).fetchone():
                    flash("This doctor profile already has a login account.", "danger")
                elif not con.execute("SELECT 1 FROM doctors WHERE id=?", (doctor_id,)).fetchone():
                    flash("Choose a valid doctor profile.", "danger")
                else:
                    con.execute("""INSERT INTO users(username,password_hash,full_name,role,doctor_id,created_at)
                        VALUES(?,?,?,?,?,?)""",
                        (username, generate_password_hash(password), full_name, "doctor", doctor_id,
                         datetime.now().isoformat(timespec="minutes")))
                    flash("Doctor account created. Share the temporary password securely and ask the doctor to change it after signing in.", "success")
                    return redirect(url_for("manage_doctors"))
    with db() as con:
        rows = con.execute("""SELECT u.id,u.username,u.full_name,u.created_at,d.name doctor_name,d.specialty
            FROM users u LEFT JOIN doctors d ON d.id=u.doctor_id WHERE u.role='doctor' ORDER BY u.full_name""").fetchall()
        available = con.execute("""SELECT d.* FROM doctors d LEFT JOIN users u ON u.doctor_id=d.id
            WHERE u.id IS NULL ORDER BY d.name""").fetchall()
    return render_template("manage_doctors.html", accounts=rows, available_doctors=available)

@app.post("/admin/doctors/<int:user_id>/delete")
@admin_required
def delete_doctor_account(user_id):
    if user_id == current_user():
        flash("You cannot delete the account currently in use.", "danger")
    else:
        with db() as con:
            con.execute("DELETE FROM users WHERE id=? AND role='doctor'", (user_id,))
        flash("Doctor login account deleted. Historical records are retained.", "warning")
    return redirect(url_for("manage_doctors"))

@app.route("/change-password", methods=["GET", "POST"])
@doctor_required
def change_password():
    if request.method == "POST":
        old_password = request.form.get("old_password", "")
        new_password = request.form.get("new_password", "")
        confirm = request.form.get("confirm_password", "")
        with db() as con:
            user = con.execute("SELECT password_hash FROM users WHERE id=?", (current_user(),)).fetchone()
            if not user or not check_password_hash(user["password_hash"], old_password):
                flash("Current password is incorrect.", "danger")
            elif len(new_password) < 10:
                flash("New password must be at least 10 characters.", "danger")
            elif new_password != confirm:
                flash("New passwords do not match.", "danger")
            else:
                con.execute("UPDATE users SET password_hash=? WHERE id=?",
                            (generate_password_hash(new_password), current_user()))
                session.clear()
                flash("Password changed. Please sign in again.", "success")
                return redirect(url_for("login"))
    return render_template("change_password.html")

@app.route("/patients", methods=["GET", "POST"])
@doctor_required
def patients():
    uid = current_user()
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        phone = request.form.get("phone", "").strip()
        try:
            age = int(request.form.get("age", ""))
        except ValueError:
            age = -1
        gender = request.form.get("gender", "")
        if not name or not phone or not gender or not (0 <= age <= 120):
            flash("Enter a valid name, age (0–120), gender and phone number.", "danger")
        else:
            with db() as con:
                con.execute("""INSERT INTO patients(name,age,gender,phone,email,address,created_at,user_id)
                    VALUES(?,?,?,?,?,?,?,?)""",
                    (name, age, gender, phone, request.form.get("email","").strip(),
                     request.form.get("address","").strip(), datetime.now().isoformat(timespec="minutes"),
                     None if session.get("role") == "admin" else uid))
            flash("Patient registered successfully.", "success")
            return redirect(url_for("patients"))
    with db() as con:
        if session.get("role") == "admin":
            rows = con.execute("SELECT * FROM patients ORDER BY id DESC").fetchall()
        else:
            rows = con.execute("SELECT * FROM patients WHERE user_id=? ORDER BY id DESC", (uid,)).fetchall()
    return render_template("patients.html", patients=rows)

@app.route("/appointments", methods=["GET", "POST"])
@doctor_required
def appointments():
    uid, role = current_user(), session.get("role")
    if request.method == "POST":
        try:
            patient_id = int(request.form["patient_id"])
            requested_doctor_id = int(request.form["doctor_id"]) if role == "admin" else int(session.get("doctor_id") or 0)
            date = request.form["appointment_date"]
            reason = request.form["reason"].strip()
            datetime.strptime(date, "%Y-%m-%dT%H:%M")
            if not reason:
                raise ValueError()
        except (ValueError, KeyError):
            flash("Please complete all appointment fields with a valid date and time.", "danger")
        else:
            with db() as con:
                if role == "admin":
                    patient = con.execute("SELECT id FROM patients WHERE id=?", (patient_id,)).fetchone()
                    owner_id = None
                else:
                    patient = con.execute("SELECT id FROM patients WHERE id=? AND user_id=?", (patient_id, uid)).fetchone()
                    owner_id = uid
                doctor = con.execute("SELECT id FROM doctors WHERE id=?", (requested_doctor_id,)).fetchone()
                if not patient or not doctor:
                    flash("Select a patient assigned to your account and a valid doctor.", "danger")
                else:
                    con.execute("""INSERT INTO appointments(patient_id,doctor_id,appointment_date,reason,status,created_at,user_id)
                        VALUES(?,?,?,?,?,?,?)""",
                        (patient_id, requested_doctor_id, date, reason, "Booked",
                         datetime.now().isoformat(timespec="minutes"), owner_id))
                    flash("Appointment booked.", "success")
                    return redirect(url_for("appointments"))
    with db() as con:
        if role == "admin":
            patients_rows = con.execute("SELECT id,name,phone FROM patients ORDER BY name").fetchall()
            doctors_rows = con.execute("SELECT * FROM doctors ORDER BY name").fetchall()
            rows = con.execute("""SELECT a.*,p.name patient_name,d.name doctor_name,d.specialty
                FROM appointments a JOIN patients p ON p.id=a.patient_id JOIN doctors d ON d.id=a.doctor_id
                ORDER BY a.appointment_date DESC""").fetchall()
        else:
            patients_rows = con.execute("SELECT id,name,phone FROM patients WHERE user_id=? ORDER BY name", (uid,)).fetchall()
            doctors_rows = con.execute("SELECT * FROM doctors WHERE id=?", (session.get("doctor_id"),)).fetchall()
            rows = con.execute("""SELECT a.*,p.name patient_name,d.name doctor_name,d.specialty
                FROM appointments a JOIN patients p ON p.id=a.patient_id JOIN doctors d ON d.id=a.doctor_id
                WHERE a.user_id=? ORDER BY a.appointment_date DESC""", (uid,)).fetchall()
    return render_template("appointments.html", patients=patients_rows, doctors=doctors_rows, appointments=rows)

@app.post("/appointments/<int:appointment_id>/status")
@doctor_required
def appointment_status(appointment_id):
    status = request.form.get("status")
    if status not in ("Booked", "Completed", "Cancelled"):
        flash("Invalid appointment status.", "danger")
    else:
        with db() as con:
            if session.get("role") == "admin":
                con.execute("UPDATE appointments SET status=? WHERE id=?", (status, appointment_id))
            else:
                con.execute("UPDATE appointments SET status=? WHERE id=? AND user_id=?",
                            (status, appointment_id, current_user()))
        flash("Appointment status updated if the record belongs to your account.", "success")
    return redirect(url_for("appointments"))

@app.route("/billing", methods=["GET", "POST"])
@doctor_required
def billing():
    uid, role = current_user(), session.get("role")
    if request.method == "POST":
        try:
            patient_id = int(request.form["patient_id"])
            amount = float(request.form["amount"])
            description = request.form["description"].strip()
            if amount < 0 or not description:
                raise ValueError()
        except (ValueError, KeyError):
            flash("Enter a valid patient, description and non-negative amount.", "danger")
        else:
            with db() as con:
                patient = con.execute("SELECT id FROM patients WHERE id=?" + ("" if role == "admin" else " AND user_id=?"),
                                      (patient_id,) if role == "admin" else (patient_id, uid)).fetchone()
                if not patient:
                    flash("Select a patient assigned to your account.", "danger")
                else:
                    con.execute("""INSERT INTO bills(patient_id,description,amount,payment_status,created_at,user_id)
                        VALUES(?,?,?,?,?,?)""",
                        (patient_id, description, amount, "Pending", datetime.now().isoformat(timespec="minutes"),
                         None if role == "admin" else uid))
                    flash("Bill created. This demo does not process real payments.", "success")
                    return redirect(url_for("billing"))
    with db() as con:
        if role == "admin":
            patients_rows = con.execute("SELECT id,name FROM patients ORDER BY name").fetchall()
            rows = con.execute("""SELECT b.*,p.name patient_name FROM bills b JOIN patients p ON p.id=b.patient_id
                ORDER BY b.id DESC""").fetchall()
        else:
            patients_rows = con.execute("SELECT id,name FROM patients WHERE user_id=? ORDER BY name", (uid,)).fetchall()
            rows = con.execute("""SELECT b.*,p.name patient_name FROM bills b JOIN patients p ON p.id=b.patient_id
                WHERE b.user_id=? ORDER BY b.id DESC""", (uid,)).fetchall()
    return render_template("billing.html", patients=patients_rows, bills=rows)

@app.post("/billing/<int:bill_id>/paid")
@doctor_required
def mark_paid(bill_id):
    with db() as con:
        if session.get("role") == "admin":
            con.execute("UPDATE bills SET payment_status='Paid' WHERE id=?", (bill_id,))
        else:
            con.execute("UPDATE bills SET payment_status='Paid' WHERE id=? AND user_id=?", (bill_id, current_user()))
    flash("Bill marked paid in this demo. Verify any real payment independently.", "warning")
    return redirect(url_for("billing"))

@app.route("/ai-assistant", methods=["GET", "POST"])
@doctor_required
def ai_assistant():
    result = None
    if request.method == "POST":
        symptoms = request.form.get("symptoms", "").lower().strip()
        if not symptoms:
            result = "Please enter the symptoms or concern."
        else:
            rules = [
                (["chest pain", "difficulty breathing", "shortness of breath", "unconscious", "stroke", "severe bleeding"],
                 "URGENT: These can be emergency warning signs. Seek emergency medical care immediately; do not rely on this demo tool."),
                (["fever", "temperature", "chills"], "General information: fever can have many causes. A clinician should assess duration, temperature, age and other symptoms."),
                (["rash", "itch", "skin"], "Possible next step: consider clinician review, especially if symptoms are spreading, painful or accompanied by fever."),
                (["cough", "sore throat", "cold"], "General information: respiratory symptoms have many causes. Seek clinical advice if severe, persistent or worsening."),
                (["headache", "migraine"], "General information: headache causes vary. Sudden severe headache, weakness, confusion or vision changes require emergency care."),
                (["stomach", "abdominal", "vomit", "diarrhea"], "General information: abdominal symptoms have many causes. Severe pain, blood, fainting or dehydration needs urgent assessment."),
            ]
            result = next((msg for keys, msg in rules if any(k in symptoms for k in keys)),
                "No rule matched. This tool cannot diagnose. Please have a qualified clinician assess the concern and patient history.")
    return render_template("ai_assistant.html", result=result)

@app.route("/export-excel")
@doctor_required
def export_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Patients"
    ws.append(["Patient ID", "Name", "Age", "Gender", "Phone", "Email", "Address", "Registered"])
    role, uid = session.get("role"), current_user()
    with db() as con:
        if role == "admin":
            patient_rows = con.execute("SELECT * FROM patients ORDER BY id").fetchall()
            appt_rows = con.execute("""SELECT a.id,p.name,d.name,d.specialty,a.appointment_date,a.reason,a.status
                FROM appointments a JOIN patients p ON p.id=a.patient_id JOIN doctors d ON d.id=a.doctor_id ORDER BY a.id""").fetchall()
            bill_rows = con.execute("""SELECT b.id,p.name,b.description,b.amount,b.payment_status,b.created_at
                FROM bills b JOIN patients p ON p.id=b.patient_id ORDER BY b.id""").fetchall()
        else:
            patient_rows = con.execute("SELECT * FROM patients WHERE user_id=? ORDER BY id", (uid,)).fetchall()
            appt_rows = con.execute("""SELECT a.id,p.name,d.name,d.specialty,a.appointment_date,a.reason,a.status
                FROM appointments a JOIN patients p ON p.id=a.patient_id JOIN doctors d ON d.id=a.doctor_id
                WHERE a.user_id=? ORDER BY a.id""", (uid,)).fetchall()
            bill_rows = con.execute("""SELECT b.id,p.name,b.description,b.amount,b.payment_status,b.created_at
                FROM bills b JOIN patients p ON p.id=b.patient_id WHERE b.user_id=? ORDER BY b.id""", (uid,)).fetchall()
        for r in patient_rows:
            ws.append([r["id"], r["name"], r["age"], r["gender"], r["phone"], r["email"], r["address"], r["created_at"]])
        wa = wb.create_sheet("Appointments")
        wa.append(["Appointment ID", "Patient", "Doctor", "Specialty", "Date & Time", "Reason", "Status"])
        for r in appt_rows:
            wa.append(list(r))
        wbills = wb.create_sheet("Bills")
        wbills.append(["Bill ID", "Patient", "Description", "Amount", "Payment Status", "Created"])
        for r in bill_rows:
            wbills.append(list(r))
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            letter = col[0].column_letter
            width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 36)
            sheet.column_dimensions[letter].width = width
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "hospital_records_admin.xlsx" if role == "admin" else f"hospital_records_{session.get('username','doctor')}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        with db() as con:
            user = con.execute("SELECT * FROM users WHERE username=? COLLATE NOCASE", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["full_name"] = user["full_name"]
            session["role"] = user["role"]
            session["doctor_id"] = user["doctor_id"]
            flash("Signed in successfully.", "success")
            return redirect(url_for("index"))
        flash("Incorrect username or password.", "danger")
    return render_template("login.html")

@app.post("/logout")
def logout():
    session.clear()
    flash("You have been logged out.", "success")
    return redirect(url_for("index"))

@app.errorhandler(404)
def not_found(_):
    return render_template("error.html", message="Page not found."), 404

@app.errorhandler(500)
def server_error(_):
    return render_template("error.html", message="Something went wrong. Check the terminal for details."), 500

init_db()

if __name__ == "__main__":
    app.run(debug=False, host="127.0.0.1", port=5000)
